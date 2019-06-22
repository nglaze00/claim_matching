"""
Created by Nicholas Glaze

Module for reordering a case patent claims document according to the US Patent & Trademark Office's
master copy. Reorders left column of an MS Word table according to the order of claims in the right column.
"""

import docx, time
from collections import defaultdict




def slice(claim):
    """
    Removes "*." from the beginning of claim
    :param claim: string
    :return: claim w/ prefix removed
    """
    return claim.split(".", 1)[1]

def compare_strings(str1, str2):
    """
    Returns the difference in word counts between str1 and str2, based on str1.
    :param str1: string
    :param str2: string
    :return: int
    """
    diff = 0
    words1 = str1.split(" ")
    words2 = str2.split(" ")
    counts1, counts2 = defaultdict(int), defaultdict(int)

    for word in words1:
        counts1[word] += 1
    for word in words2:
        counts2[word] += 1
    diff = 0
    for word in counts1.keys():
        diff += abs(counts1[word] - counts2[word])

    return diff

def compute_diffs(key, other, slice, fast):
    """
    Compares.
    :param key: list of strings
    :param other: list of strings
    :param slice: function to format strings before comparing
    :param fast: boolean; True: only computes scores for unmatched elements of other
    :return: diffs: matrix of ints where diffs[i, j] = compare_strings score between key[i] and other[j]; last index
                    in each row is index of min difference
    """
    reorder = [""] * len(other)
    matched = [False] * len(other)
    diffs = []  #Score matrix: row = scores for key
    for index, item in enumerate(key):
        diffs_row = []
        closest = (None, float("inf"))
        for index, test in enumerate(other):
            if not matched[index] or not fast:
                diff = compare_strings(slice(item), slice(test))
                diffs_row.append(diff)
                if diff < closest[1]:
                    closest = (index, diff)
            else:
                diffs_row.append(float("inf"))
        matched[closest[0]] = True
        diffs_row.append(closest[0])
        diffs.append(diffs_row)

    return diffs

def reorder_lists(key, other, diffs):
    """
    Reorders the items of other to match those of key so that string differences are minimized
    :param key: list of strs
    :param other: list of strs
    :param diffs: difference matrix computed by compute_diffs
    :return: key and other, with other reordered to match key as closely as possible
    """
    reordered = [""] * len(other)
    for i in range(len(key)):
        reordered[i] = other[diffs[i][-1]]
    return key, reordered

def table_to_lists(table):
    """
    Converts a two-column .docx table to two lists; assume blank rows between content rows
    :param table: docx table object
    :return: lists of strings, with column names as first items: left, right
    """
    left = [table.cell(0, 0).text]
    right = [table.cell(0, 1).text]
    for i in range(1, len(table.rows), 2):
        left.append(table.cell(i, 0).text)
        right.append(table.cell(i, 1).text)
    return left, right

def overwrite_table(left, right, table):
    """
    Overwrites input two-column table with left and right.
    :param left: list to populate left column
    :param right: list to populate right column
    :param table: parent table
    :return: none; modifies table
    """
    table.cell(0, 0).text, table.cell(0, 1).text = left[0], right[0]
    for i in range(1, len(left)):
        table.cell(2 * i - 1, 0).text = left[i]
        table.cell(2 * i - 1, 1).text = right[i]




def reorder_claims(file):
    """
    Reorders the claims in file, using the right column as the key.
    :param file: string - .docx filename
    :return: none; new file fil_reordered.docx created
    """
    doc = docx.Document(file)
    table = doc.tables[0]
    other, key = table_to_lists(table)
    diffs = compute_diffs(key[1:], other[1:], slice, False)
    _, reordered = reorder_lists(key[1:], other[1:], diffs)
    reordered = [other[0]] + reordered
    overwrite_table(reordered, key, table)
    docname = file.split(".")[0]
    doc.save(docname + "_reordered.docx")


try:
    reorder_claims("table.docx")
    print("Claims matched successfully. Closing in 5 seconds...")
    time.sleep(5)
except Exception as e:
    print("Claim matching failed.", e)
    time.sleep(5)

